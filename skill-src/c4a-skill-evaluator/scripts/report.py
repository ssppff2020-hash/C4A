#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""report.py — Level 4 of the C4A evaluator.

Turn the structured results into deliverables a human actually wants:
class overview, per-author detail, ranking, per-author next actions and
class-wide trends — as Markdown, CSV, Excel and a self-contained HTML
dashboard.

No third-party dependency is required except ``openpyxl`` for the optional
.xlsx export (which is skipped with a clear note when unavailable).
"""

from __future__ import annotations

import csv
import html
import json
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

from collector import FileInfo, Submission
from completeness import CompletenessResult, RoleResult, render_matrix, unmatched_files
from quality import CheckResult, CriterionResult, Score

BAR_FULL = "█"
BAR_EMPTY = "·"


# ---------------------------------------------------------------------------
# Report model
# ---------------------------------------------------------------------------

@dataclass
class AuthorReport:
    submission: Submission
    completeness: CompletenessResult
    criteria: List[CriterionResult]
    score: Score
    flags: List[dict] = field(default_factory=list)
    actions: List[dict] = field(default_factory=list)
    unmatched: List[FileInfo] = field(default_factory=list)
    llm_notes: List[str] = field(default_factory=list)
    tag: str = ""            # e.g. "SYNTHETIC FIXTURE"

    @property
    def author(self) -> str:
        return self.submission.author

    @property
    def quality_icon(self) -> str:
        ratio = self.score.quality_ratio
        if ratio >= 0.70:
            return "✅"
        if ratio >= 0.30:
            return "⚠️"
        return "❌"


@dataclass
class ClassReport:
    folder: str
    generated_at: str
    rubric_version: str
    authors: List[AuthorReport]
    non_c4_files: List[FileInfo] = field(default_factory=list)
    weights_note: str = ""
    llm_enabled: bool = False

    @property
    def ranked(self) -> List[AuthorReport]:
        return sorted(self.authors, key=lambda a: (-a.score.composite, a.author.lower()))

    def stats(self) -> Dict[str, object]:
        n = len(self.authors)
        complete = sum(1 for a in self.authors if a.completeness.band == "complete")
        partial = sum(1 for a in self.authors if a.completeness.band == "partial")
        insufficient = n - complete - partial
        avg_quality = (sum(a.score.quality_ratio for a in self.authors) / n) if n else 0.0
        avg_composite = (sum(a.score.composite for a in self.authors) / n) if n else 0.0
        flagged = sum(1 for a in self.authors if a.flags)

        # Which deliverable is most often missing, and which criterion is weakest.
        missing_counter: Dict[str, int] = {}
        for author in self.authors:
            for role in author.completeness.missing_roles():
                missing_counter[role.label_cn] = missing_counter.get(role.label_cn, 0) + 1
        most_missing = max(missing_counter.items(), key=lambda kv: kv[1]) if missing_counter else ("—", 0)

        criterion_avg: Dict[str, float] = {}
        for author in self.authors:
            for criterion in author.criteria:
                criterion_avg.setdefault(criterion.label_cn, []).append(criterion.ratio)  # type: ignore[arg-type]
        criterion_means = {k: sum(v) / len(v) for k, v in criterion_avg.items()}  # type: ignore[arg-type]
        weakest = min(criterion_means.items(), key=lambda kv: kv[1]) if criterion_means else ("—", 0.0)

        return {
            "n": n, "complete": complete, "partial": partial, "insufficient": insufficient,
            "avg_quality": avg_quality, "avg_composite": avg_composite, "flagged": flagged,
            "most_missing": most_missing, "criterion_means": criterion_means, "weakest": weakest,
            "missing_counter": missing_counter,
        }


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------

def _bar(ratio: float, width: int = 10) -> str:
    filled = int(round(ratio * width))
    return BAR_FULL * filled + BAR_EMPTY * (width - filled)


def render_markdown(report: ClassReport) -> str:
    s = report.stats()
    out: List[str] = []
    w = out.append

    w("# C4 提交自动评审报告")
    w("")
    w("> 由 `c4a-skill-evaluator`（C4A 挑战交付物）自动生成 ｜ 规则版本 %s" % report.rubric_version)
    w("")
    w("| 项目 | 值 |")
    w("|------|----|")
    w("| 生成时间 | %s |" % report.generated_at)
    w("| 扫描路径 | `%s` |" % report.folder)
    w("| 识别提交 | **%d** 位作者 |" % s["n"])
    w("| 非 C4 文件（已过滤） | %d |" % len(report.non_c4_files))
    w("| 评分权重 | %s |" % (report.weights_note or "默认（四条件等权 0.25）"))
    w("| LLM 深度复审 | %s |" % ("已启用（仅标注，不改分）" if report.llm_enabled else "未启用（纯规则，可离线复现）"))
    w("")

    # --- 一、班级总览 -------------------------------------------------------
    w("## 一、班级总览")
    w("")
    w("| 指标 | 数值 |")
    w("|------|------|")
    w("| 总提交人数 | %d |" % s["n"])
    w("| 完整提交（5/5） | %d |" % s["complete"])
    w("| 部分提交（3–4/5） | %d |" % s["partial"])
    w("| 严重缺失（<3/5） | %d |" % s["insufficient"])
    w("| 触发红牌人数 | %d |" % s["flagged"])
    w("| 平均四条件得分 | %.2f / 4.00 |" % (s["avg_quality"] * 4))
    w("| 平均综合分 | %.1f / 100 |" % s["avg_composite"])
    w("| 最常见缺口 | %s |" % _gap_phrase(s["most_missing"]))
    w("| 全班最弱维度 | %s |" % _weakest_phrase(s["weakest"]))
    w("")
    w("**四条件班级均值**")
    w("")
    w("| 维度 | 均值 | |")
    w("|------|------|---|")
    for label, value in s["criterion_means"].items():  # type: ignore[union-attr]
        w("| %s | %.2f | `%s` |" % (label, value, _bar(value)))
    w("")

    # --- 二、作者详情 -------------------------------------------------------
    w("## 二、作者详情")
    w("")
    for author in report.ranked:
        w("### %s%s" % (author.author,
                        "  `%s`" % author.tag if author.tag else ""))
        w("")
        w("版本：v%d ｜ 文件 %d 个 ｜ 完整性 %d/5 %s ｜ 四条件 %.2f/4 ｜ 综合分 **%.1f/100**"
          % (author.submission.version, len(author.submission.files),
             author.completeness.matched_count, author.completeness.icon,
             author.score.quality_4, author.score.composite))
        if len(author.submission.version_history) > 1:
            history = "、".join("v%d: %d 个文件" % (v, len(f))
                                for v, f in author.submission.version_history.items())
            w("")
            w("版本标记分布：%s —— 同一提交内出现多个版本号时，说明有文件未同步更新到最新版。"
              % history)
        w("")

        w("**完整性检查**")
        w("")
        w("| 交付物 | 状态 | 匹配文件 | 说明 |")
        w("|--------|------|----------|------|")
        for label, icon, files_str, note in render_matrix(author.completeness):
            w("| %s | %s | %s | %s |" % (label, icon, files_str, note))
        w("")

        w("**质量评审（C4 四条件）**")
        w("")
        w("| 条件 | 评级 | 得分 | 评审依据 |")
        w("|------|------|------|----------|")
        for criterion in author.criteria:
            evidence = "；".join("%s %s" % (c.icon, c.label) for c in criterion.applicable)
            w("| %s | %s | %.2f | %s |" % (
                criterion.label_cn, criterion.rating, criterion.ratio, evidence[:300]))
        w("")
        for criterion in author.criteria:
            w("<details><summary>%s 逐条证据（%s %.2f）</summary>" % (
                criterion.label_cn, criterion.rating, criterion.ratio))
            w("")
            for check in criterion.checks:
                w("- %s **%s** — %s" % (check.icon, check.label, check.detail))
            w("")
            w("</details>")
            w("")

        if author.flags:
            w("**🚩 红牌**")
            w("")
            for flag in author.flags:
                w("- **%s**（-%d 分，%s）：%s" % (
                    flag["label_cn"], flag["penalty"], flag["dimension"], flag["detail"]))
            w("")

        if author.llm_notes:
            w("**🤖 LLM 复审批注**（不参与打分）")
            w("")
            for note in author.llm_notes:
                w("- %s" % note)
            w("")

        w("**改进建议（按优先级）**")
        w("")
        if author.actions:
            for index, action in enumerate(author.actions, 1):
                w("%d. [P%d · %s] %s" % (index, action["priority"], action["area"], action["text"]))
        else:
            w("无需整改——五项齐全、四条件全部达标。")
        w("")

        if author.unmatched:
            w("<details><summary>未被归入五项交付物的文件（%d）</summary>" % len(author.unmatched))
            w("")
            for info in author.unmatched:
                w("- `%s`（%s，%s）" % (info.rel, info.ext, info.note or info.kind))
            w("")
            w("</details>")
            w("")
        w("---")
        w("")

    # --- 三、排名 -----------------------------------------------------------
    w("## 三、排名")
    w("")
    w("> 综合分 = 完整性 × %.1f + 四条件得分率 × %.1f，再 ×100。权重可用 `--weights` 调整。"
      % (0.4, 0.6))
    w("")
    w("| 排名 | 作者 | 完整性 | 四条件 | 综合分 | 红牌 |")
    w("|------|------|--------|--------|--------|------|")
    for index, author in enumerate(report.ranked, 1):
        w("| %d | %s | %d/5 %s | %.2f/4 | **%.1f** | %s |" % (
            index, author.author, author.completeness.matched_count, author.completeness.icon,
            author.score.quality_4, author.score.composite,
            "、".join(f["label_cn"] for f in author.flags) or "—"))
    w("")

    # --- 四、全班改进建议 ---------------------------------------------------
    w("## 四、全班改进建议")
    w("")
    w("1. **补齐最普遍的缺口**：%s" % _gap_phrase(s["most_missing"], long_form=True))
    w("2. **优先补强最弱维度**：%s" % _weakest_phrase(s["weakest"], long_form=True))
    w("3. **提交前自检**：把本技能指向自己的文件夹跑一遍，红牌清零再发群。")
    w("4. **补齐 AI 日志的迭代痕迹**：日志里写清「第几轮改了什么」，避免被判为「一句话指令直接提交」。")
    w("")
    if report.non_c4_files:
        w("### 附录 A：非 C4 文件（已排除）")
        w("")
        for info in report.non_c4_files[:40]:
            w("- `%s`" % info.rel)
        if len(report.non_c4_files) > 40:
            w("- …另有 %d 个" % (len(report.non_c4_files) - 40))
        w("")

    # --- 附录 B -------------------------------------------------------------
    w("### 附录 B：评测方法与已知局限")
    w("")
    w("- **纯规则、离线、确定性**：同一文件夹重复运行结果完全一致；不需要 API Key。")
    w("- **匹配是概率性的**：五项交付物采用「文件名 + 扩展名 + 内容信号」加权打分后贪心分配，")
    w("  弱命中会标为 ⚠️「需人工确认」，不会静默算作通过。")
    w("- **不做内容质量判断**：本报告评的是「是否符合 C4 的四条件」，不评价教案/代码本身写得好不好。")
    w("- **二进制文件仅按文件名匹配**：.mp4 / .png 不解析内容（只确认存在与命名）。")
    w("")
    return "\n".join(out)


def _gap_phrase(most_missing, long_form: bool = False) -> str:
    label, count = most_missing
    if not count:
        return "无 —— 所有提交的五项交付物全部齐全" if long_form else "无（五项全部齐全）"
    if long_form:
        return "%s —— 全班 %d 人未提交。" % (label, count)
    return "%s（%d 人）" % (label, count)


def _weakest_phrase(weakest, long_form: bool = False) -> str:
    label, value = weakest
    if label == "—" or value >= 0.95:
        return ("无 —— 全班四条件均达标（最弱维度 %.2f）" % value) if long_form \
            else "无（均达标）"
    if long_form:
        return "%s（均值 %.2f）。%s" % (label, value, _weakest_advice(label))
    return "%s（均值 %.2f）" % (label, value)


def _weakest_advice(label: str) -> str:
    table = {
        "可复用": "重点检查是否残留本机绝对路径、是否写了安装步骤与依赖。",
        "可执行": "重点检查是否真的交了 .skill/.py，以及代码能否 `python -m py_compile` 通过。",
        "可验证": "重点检查是否给了真实案例、预期输出与可复制的复现命令。",
        "IO 明确": "重点检查文档开头有没有一句「输入____，输出____」。",
    }
    return table.get(label, "")


# ---------------------------------------------------------------------------
# CSV / Excel
# ---------------------------------------------------------------------------

CSV_HEADER = [
    "author", "version", "files", "completeness_matched", "completeness_band",
    "quality_ratio", "quality_4", "composite", "red_flags",
    "reusable", "executable", "verifiable", "clear_io",
    "missing_deliverables", "top_action",
]


def render_rows(report: ClassReport) -> List[dict]:
    rows = []
    for author in report.ranked:
        by_key = {c.key: c for c in author.criteria}
        rows.append({
            "author": author.author,
            "version": "v%d" % author.submission.version,
            "files": len(author.submission.files),
            "completeness_matched": "%d/5" % author.completeness.matched_count,
            "completeness_band": author.completeness.label,
            "quality_ratio": round(author.score.quality_ratio, 4),
            "quality_4": round(author.score.quality_4, 3),
            "composite": round(author.score.composite, 2),
            "red_flags": "；".join(f["label_cn"] for f in author.flags),
            "reusable": round(by_key.get("reusable").ratio, 3) if by_key.get("reusable") else "",
            "executable": round(by_key.get("executable").ratio, 3) if by_key.get("executable") else "",
            "verifiable": round(by_key.get("verifiable").ratio, 3) if by_key.get("verifiable") else "",
            "clear_io": round(by_key.get("clear_io").ratio, 3) if by_key.get("clear_io") else "",
            "missing_deliverables": "；".join(r.label_cn for r in author.completeness.missing_roles()),
            "top_action": author.actions[0]["text"] if author.actions else "",
        })
    return rows


def write_csv(report: ClassReport, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=CSV_HEADER)
        writer.writeheader()
        for row in render_rows(report):
            writer.writerow(row)


def write_xlsx(report: ClassReport, path: Path) -> Optional[str]:
    """Write a two-sheet workbook. Returns a note string instead when unavailable."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return "openpyxl 未安装，已跳过 .xlsx 导出（pip install openpyxl 后重跑即可）"

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="2F5496")

    ws = wb.active
    ws.title = "Scores"
    ws.append(CSV_HEADER)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row in render_rows(report):
        ws.append([row[key] for key in CSV_HEADER])
    widths = {"author": 16, "missing_deliverables": 30, "top_action": 60, "red_flags": 24}
    for index, key in enumerate(CSV_HEADER, 1):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(key, 14)
    ws.freeze_panes = "A2"

    detail = wb.create_sheet("Detail")
    detail.append(["author", "criterion", "check", "ok", "detail"])
    for cell in detail[1]:
        cell.font = header_font
        cell.fill = PatternFill("solid", start_color="548235")
    for author in report.ranked:
        for criterion in author.criteria:
            for check in criterion.checks:
                detail.append([author.author, criterion.label_cn, check.label,
                               {True: "pass", False: "fail", None: "n/a"}[check.ok],
                               check.detail])
    detail.column_dimensions["A"].width = 16
    detail.column_dimensions["B"].width = 12
    detail.column_dimensions["C"].width = 30
    detail.column_dimensions["E"].width = 80
    detail.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return None


# ---------------------------------------------------------------------------
# Standalone HTML dashboard (no CDN, no JS framework)
# ---------------------------------------------------------------------------

_HTML_HEAD = """<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>C4 提交自动评审 · 班级仪表板</title>
<style>
 :root{--bg:#f6f7fb;--card:#fff;--ink:#1b1f2a;--muted:#6b7280;--line:#e5e7eb;
       --ok:#16a34a;--warn:#d97706;--bad:#dc2626;--accent:#2f5496}
 *{box-sizing:border-box}
 body{margin:0;padding:32px;background:var(--bg);color:var(--ink);
      font:15px/1.6 -apple-system,"Segoe UI","Microsoft YaHei",sans-serif}
 h1{font-size:26px;margin:0 0 4px} h2{font-size:19px;margin:32px 0 12px}
 .sub{color:var(--muted);margin-bottom:24px}
 .cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:12px}
 .card{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:14px 16px}
 .card .k{color:var(--muted);font-size:12px} .card .v{font-size:26px;font-weight:700;margin-top:2px}
 table{width:100%;border-collapse:collapse;background:var(--card);
       border:1px solid var(--line);border-radius:10px;overflow:hidden}
 th,td{padding:9px 12px;border-bottom:1px solid var(--line);text-align:left;vertical-align:top}
 th{background:#eef2fa;font-weight:600;font-size:13px}
 tr:last-child td{border-bottom:0}
 .num{text-align:right;font-variant-numeric:tabular-nums}
 .bar{height:8px;border-radius:4px;background:#e5e7eb;overflow:hidden;min-width:70px}
 .bar>i{display:block;height:100%}
 .ok{color:var(--ok)} .warn{color:var(--warn)} .bad{color:var(--bad)}
 .pill{display:inline-block;padding:1px 8px;border-radius:999px;font-size:12px;border:1px solid var(--line)}
 details{margin:6px 0 14px} summary{cursor:pointer;color:var(--accent)}
 .tag{background:#fff4d6;border:1px solid #f0c36d;border-radius:6px;padding:1px 6px;font-size:12px}
 code{background:#f0f2f7;padding:1px 5px;border-radius:4px;font-size:13px}
 footer{color:var(--muted);margin-top:36px;font-size:13px}
</style></head><body>
"""


def _cls(ratio: float) -> str:
    return "ok" if ratio >= 0.70 else ("warn" if ratio >= 0.30 else "bad")


def _bar_cell(ratio: float, color: str) -> str:
    return ('<div class="bar"><i style="width:%.0f%%;background:%s"></i></div>'
            % (max(2.0, ratio * 100), color))


def render_html(report: ClassReport) -> str:
    s = report.stats()
    out = [_HTML_HEAD]
    a = out.append

    a("<h1>C4 提交自动评审 · 班级仪表板</h1>")
    a('<div class="sub">扫描路径 <code>%s</code> ｜ 生成时间 %s ｜ 规则版本 %s ｜ %s</div>'
      % (html.escape(report.folder), report.generated_at, report.rubric_version,
         "规则 + LLM 批注" if report.llm_enabled else "纯规则（离线可复现）"))

    a('<div class="cards">')
    for key, value in (("提交人数", s["n"]), ("完整 5/5", s["complete"]),
                       ("部分 3–4", s["partial"]), ("严重缺失", s["insufficient"]),
                       ("触发红牌", s["flagged"]),
                       ("平均综合分", "%.1f" % s["avg_composite"])):
        a('<div class="card"><div class="k">%s</div><div class="v">%s</div></div>' % (key, value))
    a("</div>")

    a("<h2>四条件班级均值</h2><table><tr><th>维度</th><th class='num'>均值</th><th>分布</th></tr>")
    for label, value in s["criterion_means"].items():  # type: ignore[union-attr]
        a("<tr><td>%s</td><td class='num'>%.2f</td><td>%s</td></tr>"
          % (label, value, _bar_cell(value, {"ok": "#16a34a", "warn": "#d97706", "bad": "#dc2626"}[_cls(value)])))
    a("</table>")

    a("<h2>排名</h2><table><tr><th>#</th><th>作者</th><th>完整性</th><th class='num'>四条件</th>"
      "<th class='num'>综合分</th><th>红牌</th></tr>")
    for index, author in enumerate(report.ranked, 1):
        flags = "、".join(f["label_cn"] for f in author.flags) or "—"
        tag = ' <span class="tag">%s</span>' % html.escape(author.tag) if author.tag else ""
        a("<tr><td class='num'>%d</td><td><b>%s</b>%s</td><td>%d/5 %s</td>"
          "<td class='num'>%.2f</td><td class='num'><b>%.1f</b></td><td>%s</td></tr>"
          % (index, html.escape(author.author), tag,
             author.completeness.matched_count, author.completeness.icon,
             author.score.quality_4, author.score.composite, html.escape(flags)))
    a("</table>")

    a("<h2>各维度得分率</h2><table><tr><th>作者</th>"
      "<th>可复用</th><th>可执行</th><th>可验证</th><th>IO 明确</th>"
      "<th class='num'>综合分</th></tr>")
    for author in report.ranked:
        cells = []
        for criterion in author.criteria:
            ratio = criterion.ratio
            color = {"ok": "#16a34a", "warn": "#d97706", "bad": "#dc2626"}[_cls(ratio)]
            cells.append("<td>%s<div class='num' style='font-size:12px;color:#6b7280'>%.2f</div></td>"
                         % (_bar_cell(ratio, color), ratio))
        a("<tr><td><b>%s</b></td>%s<td class='num'><b>%.1f</b></td></tr>"
          % (html.escape(author.author), "".join(cells), author.score.composite))
    a("</table>")

    a("<h2>逐人详情</h2>")
    for author in report.ranked:
        a("<details><summary><b>%s</b> ｜ 完整性 %d/5 ｜ 综合分 %.1f ｜ 建议 %d 条</summary>"
          % (html.escape(author.author), author.completeness.matched_count,
             author.score.composite, len(author.actions)))
        a("<table><tr><th>交付物</th><th>状态</th><th>匹配文件</th></tr>")
        for label, icon, files_str, _note in render_matrix(author.completeness):
            a("<tr><td>%s</td><td>%s</td><td><code>%s</code></td></tr>"
              % (label, icon, html.escape(files_str)))
        a("</table><table><tr><th>条件</th><th>评级</th><th>逐条证据</th></tr>")
        for criterion in author.criteria:
            items = "<br>".join("%s %s <span style='color:#6b7280'>— %s</span>"
                                % (c.icon, html.escape(c.label), html.escape(c.detail[:160]))
                                for c in criterion.checks)
            a("<tr><td>%s</td><td>%s %.2f</td><td>%s</td></tr>"
              % (criterion.label_cn, criterion.rating, criterion.ratio, items))
        a("</table>")
        if author.actions:
            a("<ol>")
            for action in author.actions:
                a("<li>[P%d · %s] %s</li>" % (action["priority"], html.escape(action["area"]),
                                              html.escape(action["text"])))
            a("</ol>")
        a("</details>")

    a("<footer>由 c4a-skill-evaluator 生成 · C4A 挑战交付物 · "
      "综合分 = 完整性×0.4 + 四条件×0.6，满分 100</footer>")
    a("</body></html>")
    return "\n".join(out)


def write_html(report: ClassReport, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(render_html(report), encoding="utf-8")


# ---------------------------------------------------------------------------
# Machine-readable dump (for downstream tooling / the optional LLM advisor)
# ---------------------------------------------------------------------------

def to_json(report: ClassReport) -> dict:
    return {
        "folder": report.folder,
        "generated_at": report.generated_at,
        "rubric_version": report.rubric_version,
        "stats": {k: (v if not isinstance(v, tuple) else list(v))
                  for k, v in report.stats().items() if k != "criterion_means"},
        "authors": [
            {
                "author": author.author,
                "tag": author.tag,
                "version": author.submission.version,
                "completeness": {
                    "matched": author.completeness.matched_count,
                    "band": author.completeness.band,
                    "missing": [r.label_cn for r in author.completeness.missing_roles()],
                    "roles": {k: [f.rel for f in rr.files]
                              for k, rr in author.completeness.roles.items()},
                },
                "criteria": [
                    {"key": c.key, "label": c.label_cn, "ratio": round(c.ratio, 4),
                     "rating": c.rating,
                     "checks": [{"id": ck.check_id, "ok": ck.ok, "detail": ck.detail}
                                for ck in c.checks]}
                    for c in author.criteria
                ],
                "score": round(author.score.composite, 2),
                "red_flags": [f["id"] for f in author.flags],
                "actions": author.actions,
            }
            for author in report.ranked
        ],
    }


def write_json(report: ClassReport, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(to_json(report), ensure_ascii=False, indent=2), encoding="utf-8")
